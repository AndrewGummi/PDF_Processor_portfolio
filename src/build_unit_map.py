#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_unit_map.py — Утиліта зіставлення підрозділів v2
═══════════════════════════════════════════════════════
КРОК 1 — генерація таблиці зіставлення:
  python build_unit_map.py <шпо.xlsx> <залишки.xlsx> [unit_map.xlsx]

  Читає унікальні підрозділи з колонки A аркуша "1. ШПО"
  (або з docx-документів якщо передати папку замість шпо.xlsx),
  автоматично зіставляє з аркушами залишків,
  записує Excel для ручного коригування.

КРОК 2 — застосування зіставлення (перейменування аркушів у залишках):
  python build_unit_map.py --apply <unit_map.xlsx> <залишки.xlsx>

  Читає заповнену таблицю (колонка C або B),
  перейменовує аркуші у файлі залишків на точні назви підрозділів,
  щоб doc_processor знаходив їх точним збігом.
"""

import sys, os, re, io, subprocess, shutil, glob
from datetime import datetime

try:
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl",
                           "--break-system-packages", "-q"])
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

try:
    from docx import Document
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx",
                           "--break-system-packages", "-q"])
    from docx import Document

# ═══════════════════════════════════════════════════════════════════════════
# ЛОГІКА ПОШУКУ — точна копія з doc_processor.py
# ═══════════════════════════════════════════════════════════════════════════
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
def w(tag): return f"{{{W_NS}}}{tag}"

EVENT_RE = re.compile(r'внаслідок.+?(\d{2}\.\d{2}\.\d{4})', re.UNICODE | re.IGNORECASE)
UNIT_RE = re.compile(
    r'^([а-яіїєґА-ЯІЇЄҐa-zA-Z0-9][а-яіїєґА-ЯІЇЄҐa-zA-Z0-9\s\-\.]*'
    r'(?:дшб|дшр|дшбр|рвп|рбпс|рбнс|ісв|тро|мсб|мср|тб|тр|бпс|бпак|аемб|зрадн|садн|абатр|батр|вреб)'
    r'[а-яіїєґА-ЯІЇЄҐa-zA-Z0-9\s\-\.]*):?\s*$',
    re.UNICODE | re.IGNORECASE
)
NOT_UNIT_RE = re.compile(
    r'служба|начальник|командир|погоджено|прошу|відкрита|майно|медична|озброєння|логістики',
    re.IGNORECASE
)

def _sheet_key_tokens(sheet_key):
    s = re.sub(r'-?06\s*$', '', sheet_key.strip().lower())
    s = re.sub(r'\([^)]*\)', '', s)
    tokens = set(re.split(r'[\s\-_\.–]+', s))
    tokens.discard('')
    return tokens

def _doc_unit_tokens(unit_name):
    s = unit_name.strip().lower()
    s = re.sub(r'\bдшб\b', 'бат', s)
    s = re.sub(r'\bдшр\b', 'р', s)
    s = re.sub(r'\bвзв\b', 'вз', s)
    tokens = set(re.split(r'[\s\-_\.]+', s))
    tokens.discard('')
    return tokens

def find_remains_sheet(unit_name, remains):
    """remains: {sheet_key_lower: original_name}"""
    key = unit_name.strip().lower().rstrip(':')
    if key in remains:
        return key
    for sheet_key in remains:
        if sheet_key in key or key in sheet_key:
            return sheet_key
    # Роти
    rot_m  = re.search(r'(\d+)\s*дшр', key, re.IGNORECASE)
    bat_m2 = re.search(r'(\d+)\s*дшб', key, re.IGNORECASE)
    if rot_m and bat_m2:
        rot_n, bat_n2 = rot_m.group(1), bat_m2.group(1)
        for sheet_key in remains:
            sh = re.sub(r'-?06\s*$', '', sheet_key.strip().lower())
            bm = re.match(r'(\d+)\s*бат', sh)
            if not bm or bm.group(1) != bat_n2:
                continue
            after = sh.split('_', 1)[1] if '_' in sh else ''
            rm = re.search(r'(?:^|\s)(\d+)\s*р', after)
            if rm and rm.group(1) == rot_n:
                return sheet_key
        return None
    # Токенний
    doc_tokens = _doc_unit_tokens(key)
    doc_type = {t for t in doc_tokens if not t.isdigit() and t not in ('бат','р','дшб','дшр')}
    doc_bat_m = re.search(r'(\d+)\s*(?:дшб|бат)', key, re.IGNORECASE)
    doc_bat_n = doc_bat_m.group(1) if doc_bat_m else None
    best_key, best_score = None, 0
    for sheet_key in remains:
        sh_tokens = _sheet_key_tokens(sheet_key)
        if not sh_tokens:
            continue
        if doc_bat_n:
            sb = re.search(r'(\d+)\s*бат', sheet_key, re.IGNORECASE)
            if sb and sb.group(1) != doc_bat_n:
                continue
        sh_type = {t for t in sh_tokens if not t.isdigit() and t not in ('бат','р','дшб','дшр')}
        if not (doc_type & sh_type):
            continue
        score = len(doc_tokens & sh_tokens)
        if score >= 2 and score / max(len(doc_tokens), len(sh_tokens)) >= 0.4:
            if score > best_score:
                best_score, best_key = score, sheet_key
    if best_key:
        return best_key
    # Нормалізований
    for sheet_key in remains:
        sh_norm = re.sub(r'-?06\s*$', '', sheet_key.strip().lower())
        sh_norm = re.sub(r'[\s_]+', ' ', sh_norm).strip()
        parts = sh_norm.split('_', 1)
        if len(parts) == 2:
            bat_p, unit_p = parts[0].strip(), parts[1].strip()
            bm = re.search(r'(\d+)', bat_p)
            bn = bm.group(1) if bm else None
            dbm = re.search(r'(\d+)\s*(?:дшб|бат)', key)
            dbn = dbm.group(1) if dbm else None
            doc_sub = re.sub(r'\d+\s*(?:дшб|бат)', '', key).strip()
            doc_sub = re.sub(r'\s+', ' ', doc_sub).strip()
            bat_ok = (bn == dbn) if (bn and dbn) else False
            up_tok = set(re.split(r'[\s\-]+', unit_p))
            ds_tok = set(re.split(r'[\s\-]+', doc_sub))
            drm = re.search(r'^(\d+)\s*(?:р\.?|дшр)', doc_sub)
            srm = re.search(r'^(\d+)\s*(?:р\.?)', unit_p)
            if drm and srm and drm.group(1) != srm.group(1):
                continue
            sub_ok = bool(up_tok & ds_tok) or any(t in unit_p for t in ds_tok if len(t) >= 2)
            if bat_ok and sub_ok:
                return sheet_key
    return None

# ═══════════════════════════════════════════════════════════════════════════
# ДЖЕРЕЛА ПІДРОЗДІЛІВ
# ═══════════════════════════════════════════════════════════════════════════
def get_para_text(para_xml):
    return ''.join((t.text or '') for t in para_xml.iter(w('t')))

def load_units_from_shpo(shpo_excel):
    """Читає унікальні підрозділи з колонки A аркуша '1. ШПО'."""
    wb = openpyxl.load_workbook(shpo_excel, read_only=True, data_only=True)
    ws = None
    for sname in wb.sheetnames:
        if 'шпо' in sname.strip().lower():
            ws = wb[sname]
            break
    if ws is None:
        ws = wb[wb.sheetnames[0]]
        print(f"  Увага: аркуш з 'ШПО' не знайдено, використовую '{ws.title}'")
    else:
        print(f"  Читаю аркуш: '{ws.title}'")

    seen = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        val = str(row[0]).strip()
        if not val or val.lower() in ('підрозділ', ''):
            continue
        key = val.lower()
        if key not in seen:
            seen[key] = val
    wb.close()
    print(f"  ШПО: знайдено {len(seen)} унікальних підрозділів")
    return seen

def load_units_from_docs(input_dir):
    """Читає унікальні підрозділи з .docx файлів у папці."""
    units = {}
    files = []
    for pat in ['**/*.docx', '**/*.doc']:
        files.extend(glob.glob(os.path.join(input_dir, pat), recursive=True))
    files = sorted(set(files))
    print(f"  Документів знайдено: {len(files)}")
    for fpath in files:
        try:
            doc = Document(fpath)
            body = doc.element.body
            for para in body.iter(w('p')):
                text_raw = get_para_text(para).strip()
                if UNIT_RE.match(text_raw) and not NOT_UNIT_RE.search(text_raw):
                    norm = text_raw.rstrip(':').strip()
                    key = norm.lower()
                    if key not in units:
                        units[key] = norm
        except Exception as e:
            print(f"  Помилка читання {os.path.basename(fpath)}: {e}")
    print(f"  З документів: знайдено {len(units)} унікальних підрозділів")
    return units

def load_sheet_names(remains_excel):
    wb = openpyxl.load_workbook(remains_excel, read_only=True)
    names = wb.sheetnames
    wb.close()
    return names

# ═══════════════════════════════════════════════════════════════════════════
# СТИЛІ
# ═══════════════════════════════════════════════════════════════════════════
HDR_FILL    = PatternFill("solid", fgColor="1F4E79")
HDR_FONT    = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
FOUND_FILL  = PatternFill("solid", fgColor="E2EFDA")
FOUND_ALT   = PatternFill("solid", fgColor="D6EFD0")
NOTF_FILL   = PatternFill("solid", fgColor="FCE4D6")
NOTF_ALT    = PatternFill("solid", fgColor="F9D0BC")
MANUAL_FILL = PatternFill("solid", fgColor="FFFACD")
SHEETS_FILL = PatternFill("solid", fgColor="DEEAF1")
SHEETS_ALT  = PatternFill("solid", fgColor="C9DFF0")
thin   = Side(style='thin',   color='BFBFBF')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def sc(cell, fill=None, font=None, align=None):
    if fill:  cell.fill = fill
    if font:  cell.font = font
    cell.border = BORDER
    if align: cell.alignment = align

# ═══════════════════════════════════════════════════════════════════════════
# КРОК 1 — генерація таблиці зіставлення
# ═══════════════════════════════════════════════════════════════════════════
def build_map(source, remains_excel, out_path):
    """
    source: або шлях до .xlsx (ШПО), або шлях до папки з docx.
    """
    if os.path.isdir(source):
        print(f"Джерело підрозділів: папка з документами ({source})")
        units = load_units_from_docs(source)
    elif source.lower().endswith(('.xlsx', '.xlsm', '.xls')):
        print(f"Джерело підрозділів: ШПО ({source})")
        units = load_units_from_shpo(source)
    else:
        print(f"Помилка: '{source}' — не папка і не Excel-файл.")
        sys.exit(1)

    print(f"Читаємо аркуші залишків: {remains_excel}")
    sheet_names  = load_sheet_names(remains_excel)
    remains_keys = {s.strip().lower(): s for s in sheet_names}
    print(f"  Аркушів: {len(sheet_names)}")

    wb = openpyxl.Workbook()

    # ── Аркуш «Зіставлення» ─────────────────────────────────────────────
    ws = wb.active
    ws.title = "Зіставлення"

    headers = [
        "Підрозділ (зі штату / документів)",
        "Аркуш залишків (знайдено авто)",
        "← Виправ якщо авто неправильно або порожньо",
        "Статус",
    ]
    col_w = [44, 44, 48, 16]
    for col, (h, cw) in enumerate(zip(headers, col_w), 1):
        cell = ws.cell(row=1, column=col, value=h)
        sc(cell, fill=HDR_FILL, font=HDR_FONT,
           align=Alignment(horizontal='center', vertical='center', wrap_text=True))
        ws.column_dimensions[get_column_letter(col)].width = cw
    ws.row_dimensions[1].height = 36
    ws.freeze_panes = "A2"

    # Незнайдені першими, потім за алфавітом
    sorted_units = sorted(
        units.items(),
        key=lambda kv: (1 if find_remains_sheet(kv[0], remains_keys) else 0, kv[0])
    )

    found_n = not_found_n = 0
    for row_i, (key, orig) in enumerate(sorted_units, 2):
        auto_key  = find_remains_sheet(key, remains_keys)
        auto_orig = remains_keys[auto_key] if auto_key else ""

        if auto_key:
            status   = "✓ знайдено"
            row_fill = FOUND_FILL if row_i % 2 == 0 else FOUND_ALT
            st_color = "375623"
            found_n += 1
        else:
            status      = "⚠ не знайдено"
            row_fill    = NOTF_FILL if row_i % 2 == 0 else NOTF_ALT
            st_color    = "C00000"
            not_found_n += 1

        for col, val in enumerate([orig, auto_orig, "", status], 1):
            cell = ws.cell(row=row_i, column=col, value=val)
            fill = MANUAL_FILL if col == 3 else row_fill
            fnt  = Font(name="Calibri", size=10,
                        color=st_color if col == 4 else "000000")
            sc(cell, fill=fill, font=fnt,
               align=Alignment(vertical='center', wrap_text=(col in (1,2,3))))
        ws.row_dimensions[row_i].height = 18

    # Підсумок
    sr = len(sorted_units) + 2
    ws.cell(row=sr, column=1,
            value=f"Всього: {len(units)}   |   Знайдено: {found_n}   |   "
                  f"Потребують ручного заповнення: {not_found_n}")
    ws.cell(row=sr, column=1).font = Font(bold=True, name="Calibri", size=10, color="1F4E79")
    ws.merge_cells(start_row=sr, start_column=1, end_row=sr, end_column=4)

    # Інструкція
    ir = sr + 2
    instr = (
        "Інструкція:\n"
        "1. Для рядків де колонка B порожня або неправильна — вкажи точну назву аркуша "
        "у колонці C (скопіюй з аркуша «Аркуші залишків»).\n"
        "2. Після заповнення запусти:\n"
        "     python build_unit_map.py --apply unit_map.xlsx залишки.xlsx\n"
        "   Це перейменує аркуші залишків на точні назви підрозділів — і надалі матч буде точним.\n"
        "3. Потім запускай doc_processor.py як зазвичай (--unit-map більше не потрібен).\n\n"
        "   АБО: передай цей файл у doc_processor.py через --unit-map без перейменування."
    )
    ws.cell(row=ir, column=1, value=instr)
    ws.cell(row=ir, column=1).font = Font(italic=True, name="Calibri", size=9, color="595959")
    ws.cell(row=ir, column=1).alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(start_row=ir, start_column=1, end_row=ir+4, end_column=4)
    ws.row_dimensions[ir].height = 90

    # ── Аркуш «Аркуші залишків» ──────────────────────────────────────────
    ws2 = wb.create_sheet("Аркуші залишків")
    ws2.column_dimensions["A"].width = 52
    ws2.column_dimensions["B"].width = 22

    for col, h in enumerate(["Назва аркуша (таблиця залишків)", "Підрозділів зі штату"], 1):
        cell = ws2.cell(row=1, column=col, value=h)
        sc(cell, fill=HDR_FILL, font=HDR_FONT,
           align=Alignment(horizontal='center', vertical='center'))
    ws2.row_dimensions[1].height = 26
    ws2.freeze_panes = "A2"

    usage = {k: [] for k in remains_keys}
    for key, orig in units.items():
        found = find_remains_sheet(key, remains_keys)
        if found and found in usage:
            usage[found].append(orig)

    for row_i, sh_name in enumerate(sheet_names, 2):
        k = sh_name.strip().lower()
        u = usage.get(k, [])
        fill = (FOUND_FILL if row_i % 2 == 0 else FOUND_ALT) if u else \
               (SHEETS_FILL if row_i % 2 == 0 else SHEETS_ALT)
        c1 = ws2.cell(row=row_i, column=1, value=sh_name)
        c2 = ws2.cell(row=row_i, column=2, value=len(u) if u else "")
        for cell in (c1, c2):
            sc(cell, fill=fill, font=Font(name="Calibri", size=10),
               align=Alignment(vertical='center'))
        ws2.row_dimensions[row_i].height = 16

    wb.save(out_path)
    print(f"\nФайл зіставлення збережено: {out_path}")
    print(f"  Знайдено автоматично:          {found_n}/{len(units)}")
    print(f"  Потребують ручного заповнення: {not_found_n}")
    print(f"\nДалі:")
    print(f"  1. Відкрий {out_path}, заповни колонку C для ⚠-рядків")
    print(f"  2. python build_unit_map.py --apply {out_path} {remains_excel}")

# ═══════════════════════════════════════════════════════════════════════════
# КРОК 2 — перейменування аркушів у файлі залишків
# ═══════════════════════════════════════════════════════════════════════════
def apply_map(map_excel, remains_excel):
    """
    Читає заповнену таблицю зіставлення.
    Перейменовує аркуші у файлі залишків на точні назви підрозділів.
    Резервна копія файлу залишків зберігається автоматично.
    """
    print(f"Читаємо таблицю зіставлення: {map_excel}")
    wb_map = openpyxl.load_workbook(map_excel, read_only=True, data_only=True)
    ws_map = None
    for sname in wb_map.sheetnames:
        if 'зіставлення' in sname.lower():
            ws_map = wb_map[sname]
            break
    if ws_map is None:
        ws_map = wb_map[wb_map.sheetnames[0]]

    # {old_sheet_name_lower: (old_original, new_name, unit_original)}
    renames = {}
    skipped = []

    for row in ws_map.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        unit_name  = str(row[0]).strip() if row[0] else ""
        auto_sheet = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        manual     = str(row[2]).strip() if len(row) > 2 and row[2] else ""

        if not unit_name or unit_name.startswith("Всього"):
            continue

        target_sheet = manual or auto_sheet
        if not target_sheet:
            skipped.append(unit_name)
            continue

        # Нова назва = підрозділ з ШПО (≤31 символ, без заборонених символів)
        new_name = unit_name[:31]
        new_name = re.sub(r'[\[\]*\\/\?:]', '-', new_name).strip()

        old_key = target_sheet.lower()
        # Якщо декілька підрозділів вказують на один аркуш — беремо перший
        if old_key not in renames:
            renames[old_key] = (target_sheet, new_name, unit_name)

    wb_map.close()

    if not renames:
        print("Немає записів для перейменування.")
        return

    print(f"  Записів для перейменування: {len(renames)}")
    if skipped:
        print(f"  Пропущено (немає аркуша): {len(skipped)}")
        for s in skipped:
            print(f"    - {s}")

    # Резервна копія
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    base, ext = os.path.splitext(remains_excel)
    backup = f"{base}_backup_{ts}{ext}"
    shutil.copy2(remains_excel, backup)
    print(f"  Резервна копія: {backup}")

    wb = openpyxl.load_workbook(remains_excel)
    renamed = []
    not_found = []

    for sheet in wb.worksheets:
        old_key = sheet.title.strip().lower()
        if old_key in renames:
            old_orig, new_name, unit_orig = renames[old_key]
            if sheet.title != new_name:
                print(f"  '{sheet.title}'  →  '{new_name}'")
                sheet.title = new_name
                renamed.append((old_orig, new_name))
            else:
                print(f"  '{sheet.title}' — вже правильна назва, пропускаємо")
                renamed.append((old_orig, new_name))

    # Які аркуші не знайшли
    renamed_old_lower = {r[0].lower() for r in renamed}
    all_sheet_lower   = {s.title.lower() for s in wb.worksheets}
    for old_key, (old_orig, new_name, unit_orig) in renames.items():
        if old_orig.lower() not in renamed_old_lower and old_key not in all_sheet_lower:
            not_found.append(old_orig)

    wb.save(remains_excel)
    wb.close()

    print(f"\nГотово!")
    print(f"  Перейменовано аркушів: {len(renamed)}")
    if not_found:
        print(f"  Не знайдено в файлі залишків ({len(not_found)}):")
        for n in not_found:
            print(f"    - {n}")
    print(f"\nТепер doc_processor знайде підрозділи точним збігом без --unit-map.")
    print(f"(Резервна копія: {backup})")

# ═══════════════════════════════════════════════════════════════════════════
# MAIN
# ═══════════════════════════════════════════════════════════════════════════
def main():
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')

    args = sys.argv[1:]

    if args and args[0] == '--apply':
        # Крок 2: python build_unit_map.py --apply unit_map.xlsx залишки.xlsx
        if len(args) < 3:
            print("Використання: build_unit_map.py --apply <unit_map.xlsx> <залишки.xlsx>")
            sys.exit(1)
        apply_map(args[1], args[2])
    else:
        # Крок 1: python build_unit_map.py <шпо.xlsx або папка> <залишки.xlsx> [unit_map.xlsx]
        if len(args) < 2:
            print("Використання:")
            print("  build_unit_map.py <шпо.xlsx>    <залишки.xlsx> [unit_map.xlsx]")
            print("  build_unit_map.py <папка_docx>  <залишки.xlsx> [unit_map.xlsx]")
            print("  build_unit_map.py --apply <unit_map.xlsx> <залишки.xlsx>")
            sys.exit(1)
        source        = args[0]
        remains_excel = args[1]
        out_path      = args[2] if len(args) > 2 else \
                        os.path.join(os.path.dirname(os.path.abspath(remains_excel)), "unit_map.xlsx")
        build_map(source, remains_excel, out_path)

if __name__ == "__main__":
    main()
