#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
import_pidrozdily_xlsm.py

Імпортує щомісячний файл "Залишки_На_ХХ_20ХХ.xlsm" (один аркуш = один
підрозділ, стовпці = місяці року) в аркуш "Підрозділи" майстер-файлу
обліку майна бригади, і перебудовує "Залишки — Загалом" / "Звірка"
під актуальний список найменувань.

ВАЖЛИВО: тут НЕМАЄ ані Tesseract, ані vision-LLM, ані виклику API. Це
джерело -- вже типізований, структурований Excel-файл (не фото, не
скан), тож дані читаються напряму й на 100% точно через openpyxl.
Vision-LLM (cards_photo_to_excel.py) потрібен лише для рукописних
паперових карток; сюди його підключати не треба й не варто.

Кожен запуск ПОВНІСТЮ ЗАМІНЮЄ вміст аркуша "Підрозділи" свіжим знімком
(а не дописує) -- бо це стан "зараз", а не журнал операцій. Якщо
потрібна історія по місяцях -- запускайте на кожен місячний файл окремо
і зберігайте копії master-файлу (наприклад Облік_майна_бригади_2026-06.xlsx),
а не перезаписуйте один і той самий щомісяця.

ВИКОРИСТАННЯ (Git Bash):
    python import_pidrozdily_xlsm.py \\
        --input ./Залишки_На_07_2026.xlsm \\
        --output ./Облік_майна_бригади.xlsx
"""

import argparse
import os
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

import oblik_common as oc

SHEET_NAME = "Підрозділи"
TABLE_NAME = "Pidrozdily"
HEADERS = ["№", "Назва підрозділу", "Найменування майна", "Одиниця виміру",
           "Кількість", "Дата обліку", "Відповідальна особа", "Телефон", "Примітка"]


def guess_year(input_path: str) -> int:
    m = re.search(r"20\d{2}", Path(input_path).stem)
    if m:
        return int(m.group(0))
    return datetime.today().year


def replace_pidrozdily_sheet(wb, records, source_name: str):
    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]
    ws = wb.create_sheet(SHEET_NAME)
    oc.style_header(ws, HEADERS)

    for i, (subdivision, item_name, unit, qty, record_date) in enumerate(records, start=2):
        ws.cell(row=i, column=1, value=f"=ROW()-1").font = Font(name="Arial")
        ws.cell(row=i, column=2, value=subdivision).font = Font(name="Arial")
        ws.cell(row=i, column=3, value=item_name).font = Font(name="Arial")
        ws.cell(row=i, column=4, value=unit).font = Font(name="Arial")
        c_qty = ws.cell(row=i, column=5, value=qty)
        c_qty.font = Font(name="Arial")
        c_qty.number_format = "#,##0"
        c_date = ws.cell(row=i, column=6, value=record_date)
        c_date.font = Font(name="Arial")
        c_date.number_format = "DD.MM.YYYY"
        # колонки 7-8 (Відповідальна особа / Телефон) -- формули, див. stamp_responsible_lookup нижче
        ws.cell(row=i, column=9, value=f"з {source_name}").font = Font(name="Arial", size=9, italic=True)

    oc.add_or_resize_table(ws, TABLE_NAME, len(HEADERS), len(records) + 1)
    oc.stamp_responsible_lookup(ws, TABLE_NAME, HEADERS, "Назва підрозділу")
    for col_idx, w in zip(range(1, 10), [6, 26, 45, 12, 12, 14, 24, 16, 22]):
        ws.column_dimensions[get_column_letter(col_idx)].width = w
    return ws


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, help="Файл 'Залишки_На_ХХ_20ХХ.xlsm'")
    parser.add_argument("--output", required=True, help="Майстер-файл обліку (.xlsx)")
    parser.add_argument("--year", type=int, default=None,
                         help="Рік даних (за замовчуванням береться з імені файлу, інакше поточний)")
    args = parser.parse_args()

    input_path = os.path.expanduser(args.input)
    output_path = Path(os.path.expanduser(args.output))
    year = args.year or guess_year(input_path)

    print(f"Джерело: {input_path}")
    print(f"Рік даних: {year}")

    wb_source = load_workbook(input_path, data_only=True)
    records = oc.parse_subdivisions_xlsm(wb_source, year)
    print(f"Знайдено рядків майна по підрозділах: {len(records)} "
          f"({len(wb_source.sheetnames)} аркушів/підрозділів)")

    if output_path.exists():
        wb = load_workbook(output_path)
    else:
        wb = Workbook()
        wb.remove(wb.active)

    replace_pidrozdily_sheet(wb, records, Path(input_path).name)
    oc.ensure_responsible_sheet(wb)
    n_items = oc.rebuild_rollup(wb)
    print(f"Унікальних найменувань майна в зведенні (Журнал+Склади+Підрозділи): {n_items}")

    if "Журнал руху" not in wb.sheetnames:
        print("УВАГА: аркуша 'Журнал руху' ще нема -- запустіть спочатку "
              "cards_photo_to_excel.py або build_master_workbook.py.")

    wb.save(output_path)
    print(f"Готово: {output_path}")


if __name__ == "__main__":
    main()
