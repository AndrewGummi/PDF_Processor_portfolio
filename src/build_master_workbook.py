#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_master_workbook.py -- створює структуру майстер-файлу обліку майна
бригади З НУЛЯ:

    Журнал руху          -- SYNTHETIC demo rows (invented items, document
                             numbers and quantities); uncertain fields are
                             flagged "Потребує перевірки"
    Склади               -- ПРИКЛАД-рядок (реальних даних по складах ще нема)
    Залишки — Загалом    -- зведення по бригаді, SUMIFS з "Журналу руху"
    Звірка               -- порівняння з (Склади + Підрозділи)

Аркуш "Підрозділи" тут НЕ створюється -- для нього є окремий
import_pidrozdily_xlsm.py, який імпортує дані з
"Залишки_На_ХХ_20ХХ.xlsm" (ніякого OCR/LLM не треба, файл вже
типізований Excel). Запускайте build_master_workbook.py один
раз для початкової структури, а тоді import_pidrozdily_xlsm.py.
"""

from pathlib import Path
import datetime

import cards_photo_to_excel as cpe
import oblik_common as oc
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

OUTPUT = Path("output_excel/Облік_майна_бригади.xlsx")


def build_journal(wb):
    """Builds the "Журнал руху" (movement journal) sheet from SYNTHETIC demo
    rows, showing the layout and how an uncertain reading is flagged
    ("Потребує перевірки" + yellow fill, exactly as cards_photo_to_excel.py
    marks it on real input)."""
    ws = wb.active
    ws.title = cpe.JOURNAL_SHEET_NAME
    cpe._write_journal_header(ws)

    # Invented item names, document numbers, dates and quantities. The point
    # is to demonstrate the journal structure and the review flag, not to
    # ship any real record.
    cards = [
        ("Sample item A", [{
            "record_date": "12.06.2024", "doc_name": "Акт", "doc_number": "01/001",
            "doc_date": "25.04.2024", "supplier": None, "received": "100",
            "issued": None, "stock_total": "100", "stock_cat1": "100",
            "review_needed": True,
            "notes": "Supplier unreadable; right-hand columns cut off by the frame edge"
        }]),
        ("Sample item B", [{
            "record_date": "12.06.2024", "doc_name": "Акт", "doc_number": "01/001",
            "doc_date": "25.04.2024", "supplier": None, "received": "50",
            "issued": None, "stock_total": "50", "stock_cat1": "50",
            "review_needed": True,
            "notes": "Supplier unreadable; right-hand columns cut off by the frame edge"
        }]),
        ("Sample item C", [
            {"record_date": "12.06.2024", "doc_name": "Акт", "doc_number": "01/001",
             "doc_date": "25.04.2024", "supplier": None, "received": None,
             "issued": None, "stock_total": None,
             "review_needed": True,
             "notes": "Quantity illegible -- deliberately NOT guessed, check the original"},
            {"record_date": "22.06.2024", "doc_name": "Роздавальна відомість",
             "doc_number": "02/002",
             "doc_date": "15.06.2024", "supplier": "Sample supplier", "received": None,
             "issued": "200", "stock_total": "800", "stock_cat1": "800",
             "review_needed": True,
             "notes": "Record date uncertain; right-hand columns cut off"},
            {"record_date": "23.06.2024", "doc_name": "Роздавальна відомість",
             "doc_number": "02/003",
             "doc_date": "15.06.2024", "supplier": "Sample supplier", "received": None,
             "issued": "100", "stock_total": "700", "stock_cat1": "700",
             "review_needed": True,
             "notes": "Record date uncertain; right-hand columns cut off"},
        ]),
        ("Sample item D", [{
            "record_date": "12.06.2024", "doc_name": "Акт", "doc_number": "01/001",
            "doc_date": "25.04.2024", "supplier": None, "received": "300",
            "issued": None, "stock_total": "300", "stock_cat1": "300",
            "review_needed": True,
            "notes": "Supplier unreadable; right-hand columns cut off by the frame edge"
        }]),
        ("Sample item E", [{
            "record_date": "12.06.2024", "doc_name": "Акт", "doc_number": "01/001",
            "doc_date": "25.04.2024", "supplier": None, "received": "400",
            "issued": None, "stock_total": "400", "stock_cat1": "400",
            "review_needed": True,
            "notes": "Supplier unreadable; right-hand columns cut off by the frame edge"
        }]),
    ]

    for item_name, rows in cards:
        cpe.append_rows(ws, {"item_name": item_name, "rows": rows}, "sample card (demo)")

    cpe.autosize_columns(ws)
    cpe.register_journal_table(ws)
    return ws


def build_sklady_example(wb):
    ws = wb.create_sheet("Склади")
    headers = ["№", "Назва складу", "Найменування майна", "Одиниця виміру",
               "Кількість", "Дата обліку", "Відповідальна особа", "Телефон", "Примітка"]
    oc.style_header(ws, headers)
    row = ["=ROW()-1", "Приклад: Склад №1 (речовий)", "Термос", "шт", 50,
           datetime.date.today(), None, None,
           "ПРИКЛАД -- замініть на реальні дані по ваших складах"]
    for col, value in enumerate(row, start=1):
        c = ws.cell(row=2, column=col, value=value)
        c.font = Font(name="Arial")
        c.fill = oc.EXAMPLE_FILL
        if col == 5:
            c.number_format = "#,##0"
        if col == 6:
            c.number_format = "DD.MM.YYYY"
    oc.add_or_resize_table(ws, "Sklady", len(headers), 2)
    oc.stamp_responsible_lookup(ws, "Sklady", headers, "Назва складу")
    for col_idx, w in zip(range(1, 10), [6, 24, 30, 12, 12, 14, 24, 16, 40]):
        ws.column_dimensions[get_column_letter(col_idx)].width = w
    return ws


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    build_journal(wb)
    build_sklady_example(wb)
    oc.ensure_responsible_sheet(wb)
    oc.rebuild_rollup(wb)  # "Підрозділи" ще нема -- звірка поки що без неї
    wb.save(OUTPUT)
    print(f"Готово: {OUTPUT}")
    print("Далі: python import_pidrozdily_xlsm.py --input <Залишки_На_ХХ_20ХХ.xlsm> "
          f"--output {OUTPUT}")


if __name__ == "__main__":
    main()
