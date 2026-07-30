#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
import_balances_docx.py

Імпортує ОДИН docx-файл формату "ВІДОМІСТЬ наявності речового майна <Х>"
(той самий формат, що й один аркуш "Залишки_На_ХХ_20ХХ.xlsm", тільки по
одному складу/підрозділу за раз) у майстер-файл обліку. На відміну від
import_pidrozdily_xlsm.py (який ПОВНІСТЮ замінює аркуш "Підрозділи" --
бо xlsm містить усі підрозділи одразу), цей скрипт чіпає ЛИШЕ рядки
ОДНІЄЇ локації (--type sklad/pidrozdil), решту таблиці не чіпає.

Жодного OCR/LLM -- docx вже друкований/типізований, читається напряму
через python-docx.

ВИКОРИСТАННЯ (Git Bash):
    python import_balances_docx.py \\
        --input "Залишки__батрАР_-_06.docx" \\
        --output "Облік_майна_бригади.xlsx" \\
        --type pidrozdil
"""

import argparse
import os
from pathlib import Path

from openpyxl import Workbook, load_workbook

import oblik_common as oc

LOCATION_CONFIG = {
    "sklad": {
        "sheet": "Склади",
        "table": "Sklady",
        "location_col": "Назва складу",
        "headers": ["№", "Назва складу", "Найменування майна", "Одиниця виміру",
                    "Кількість", "Дата обліку", "Відповідальна особа", "Телефон", "Примітка"],
    },
    "pidrozdil": {
        "sheet": "Підрозділи",
        "table": "Pidrozdily",
        "location_col": "Назва підрозділу",
        "headers": ["№", "Назва підрозділу", "Найменування майна", "Одиниця виміру",
                    "Кількість", "Дата обліку", "Відповідальна особа", "Телефон", "Примітка"],
    },
}


def ensure_sheet(wb, cfg):
    if cfg["sheet"] in wb.sheetnames:
        return wb[cfg["sheet"]]
    ws = wb.create_sheet(cfg["sheet"])
    oc.style_header(ws, cfg["headers"])
    oc.add_or_resize_table(ws, cfg["table"], len(cfg["headers"]), 2)
    return ws


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, help="Один docx-файл 'ВІДОМІСТЬ...'")
    parser.add_argument("--output", required=True, help="Майстер-файл обліку (.xlsx)")
    parser.add_argument("--type", choices=["sklad", "pidrozdil"], default="pidrozdil",
                         help="Це склад чи підрозділ? (за замовчуванням: підрозділ)")
    parser.add_argument("--location-name", default=None,
                         help="Назва складу/підрозділу (за замовчуванням береться з тексту docx)")
    parser.add_argument("--year", type=int, default=None,
                         help="Рік даних (за замовчуванням поточний)")
    args = parser.parse_args()

    import datetime
    year = args.year or datetime.date.today().year
    input_path = os.path.expanduser(args.input)
    output_path = Path(os.path.expanduser(args.output))

    location_name, records = oc.parse_vidomist_docx(input_path, year, args.location_name)
    print(f"Локація: {location_name}")
    print(f"Знайдено позицій майна: {len(records)}")

    cfg = LOCATION_CONFIG[args.type]

    if output_path.exists():
        wb = load_workbook(output_path)
    else:
        wb = Workbook()
        wb.remove(wb.active)

    ws = ensure_sheet(wb, cfg)
    source_note = Path(input_path).name
    new_data_rows = [(item, unit, qty, rdate, f"з {source_note}")
                      for item, unit, qty, rdate in records]
    n_rows = oc.upsert_location_rows(ws, cfg["table"], cfg["headers"], cfg["location_col"],
                                      location_name, new_data_rows)
    oc.stamp_responsible_lookup(ws, cfg["table"], cfg["headers"], cfg["location_col"])
    print(f"Рядків у таблиці '{cfg['sheet']}' після оновлення: {n_rows}")

    oc.ensure_responsible_sheet(wb)
    n_items = oc.rebuild_rollup(wb)
    print(f"Унікальних найменувань майна в зведенні: {n_items}")

    wb.save(output_path)
    print(f"Готово: {output_path}")


if __name__ == "__main__":
    main()
